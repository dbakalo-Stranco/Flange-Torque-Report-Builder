"""Builds a standalone "corrections log" spreadsheet from the QC notes
collected on a batch of report entries -- what the tag showed vs. what got
corrected and why -- so it can be checked against the physical tags.
Per standing instruction, these notes never get written onto the report
itself (see report_builder.py); this is the only place they end up.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER = ["Tag #", "Drawing / P&ID#", "Flange #", "Note"]


def build_corrections_log(entries):
    """entries: the same list of dicts passed to build_workbook(). Only
    entries carrying a non-empty 'comment' produce a row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tag Corrections"
    ws.append(HEADER)

    rows = 0
    for e in entries:
        comment = (e.get("comment") or "").strip()
        if not comment:
            continue
        ws.append([e.get("tag", ""), e.get("pid", ""), e.get("flange_no", ""), comment])
        rows += 1

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    widths = [20, 26, 12, 70]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row in range(2, ws.max_row + 1):
        for col in range(1, len(HEADER) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), rows
