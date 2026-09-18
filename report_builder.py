"""
Builds Stranco Flange Torque Report workbooks/PDFs from a template .xltx.

This mirrors the cell mapping worked out by hand for Daniel Bakalo's crew:
  B4  = System Name (always 'PH02')
  G4  = Date
  B6  = P&ID #
  B8  = Flange #
  E8  = Tag #
  B16 = Size
  D16 = Class
  F16 = Bolt OD  (numeric fraction on most sheets, literal text string on
                  sheets whose F16 cell is formatted as text -- e.g. 20"/24")
  G18 = 'X' marks Hand Torque method
  B24..B27 = the 4 breakout-load stages (25/50/75/100%)
  A31 (4-bolt sheet) or B31 (all others) = Torqued By
  F39/H39/K39 = Tool label / serial / cal date
  F32..K40 = "Tools Required" grid -- separate preprinted row per tool
             capacity (32 TORQUE PUMP, 33 250LB CLICKER, 34 600LB CLICKER,
             35 W2000, 36 W4000, 37 W8000, 38 S3000, 39 80LB CLICKER/B-RAD
             JGUN, 40 open/custom row). The entry's tool_row picks which
             row gets the serial/cal date; every other row is cleared so
             no leftover template sample data bleeds through.
  Comments (A35) are intentionally never written by this module -- QC
             flags/overrides belong in a separate corrections log, not on
             the report itself (standing instruction).
  Stranco Rep block: A57 = Print Name, G57 = Signature (Michael Bakalo's
             actual signature, embedded as an image), J57 = Date -- filled
             in on every report per standing instruction, until told
             otherwise. QA/QC Rep block (A59,G59,J59) is always left blank
             for a wet signature.

Sheet selection is by bolt count, matching the tab naming already baked
into the template ('4 Bolt', '8 Bolt', '12 Bolt ' -- note trailing space,
'16 Bolt ', '20 Bolt', '24 Bolt'). openpyxl's copy_worksheet() does not
carry images over, so every master sheet's images are captured up front
and manually re-attached to each generated sheet.
"""
import io
import os
import subprocess
import tempfile
import uuid
from copy import copy

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import column_index_from_string

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "template.xltx")

# Standing rule (per Daniel, until told otherwise): auto-fill the Stranco
# Rep block with this name + his actual signature image on every generated
# report. The QA/QC Rep block is left blank (wet signature), unchanged.
STRANCO_REP_NAME = "Michael Bakalo"
SIGNATURE_IMAGE_PATH = os.path.join(os.path.dirname(__file__), "michael_bakalo_signature.png")
SIGNATURE_IMAGE_WIDTH_PX = 150
SIGNATURE_IMAGE_HEIGHT_PX = int(SIGNATURE_IMAGE_WIDTH_PX * 230 / 837)

SHEET_BY_BOLTCOUNT = {
    4: "4 Bolt",
    8: "8 Bolt",
    12: "12 Bolt ",
    16: "16 Bolt ",
    20: "20 Bolt",
    24: "24 Bolt",
}

MASTER_SHEETS = ["Example", "4 Bolt", "8 Bolt", "12 Bolt ", "16 Bolt ", "20 Bolt", "24 Bolt"]


def _load_master(template_path=TEMPLATE_PATH):
    wb = openpyxl.load_workbook(template_path, data_only=False)
    imgs = {}
    for boltcount, sheet_name in SHEET_BY_BOLTCOUNT.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        pics = [(img.ref.getvalue(), img.anchor) for img in ws._images]
        if pics:
            imgs[boltcount] = pics

    # The '4 Bolt' master ships with no embedded images in this template.
    # Borrow the 4-bolt diagram from 'Example' and the logo/anchor slot
    # from '8 Bolt' (same logo used on every sheet).
    if 4 not in imgs or not imgs[4]:
        ws_ex = wb["Example"] if "Example" in wb.sheetnames else None
        ws_8 = wb["8 Bolt"] if "8 Bolt" in wb.sheetnames else None
        diagram_data = None
        if ws_ex is not None:
            for img in ws_ex._images:
                d = img.ref.getvalue()
                if len(d) == 22661:  # known size of the 4-bolt diagram asset
                    diagram_data = d
        logo_data = None
        logo_anchor = None
        diagram_anchor4 = None
        if ws_8 is not None:
            for img in ws_8._images:
                d = img.ref.getvalue()
                if len(d) == 80427:  # known size of the Stranco logo asset
                    logo_data = d
                    logo_anchor = img.anchor
                elif len(d) == 46445:
                    diagram_anchor4 = img.anchor
        pics = []
        if logo_data is not None:
            pics.append((logo_data, logo_anchor))
        if diagram_data is not None and diagram_anchor4 is not None:
            pics.append((diagram_data, diagram_anchor4))
        if pics:
            imgs[4] = pics

    return wb, imgs


def _add_images(ws, imgs, boltcount):
    for data, anchor in imgs.get(boltcount, []):
        if data is None or anchor is None:
            continue
        img = XLImage(io.BytesIO(data))
        img.anchor = copy(anchor)
        ws.add_image(img)


def _add_signature_image(ws, cell_ref="G57", col_off_px=-6, row_off_px=-4):
    """Drops Michael Bakalo's actual (scanned, background-removed) signature
    into the Stranco Rep signature cell as an image."""
    if not os.path.exists(SIGNATURE_IMAGE_PATH):
        return
    EMU_PER_PX = 9525
    col_letter = "".join(ch for ch in cell_ref if ch.isalpha())
    row_num = int("".join(ch for ch in cell_ref if ch.isdigit()))
    col_idx = column_index_from_string(col_letter) - 1
    row_idx = row_num - 1

    img = XLImage(SIGNATURE_IMAGE_PATH)
    img.width = SIGNATURE_IMAGE_WIDTH_PX
    img.height = SIGNATURE_IMAGE_HEIGHT_PX
    marker = AnchorMarker(
        col=col_idx, colOff=col_off_px * EMU_PER_PX,
        row=row_idx, rowOff=row_off_px * EMU_PER_PX,
    )
    size = XDRPositiveSize2D(
        cx=SIGNATURE_IMAGE_WIDTH_PX * EMU_PER_PX,
        cy=SIGNATURE_IMAGE_HEIGHT_PX * EMU_PER_PX,
    )
    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


def _bolt_od_value(ws, cell_ref, bolt_od):
    """Write bolt OD respecting whichever convention (numeric fraction vs
    literal text) the target sheet's F16 cell already uses."""
    fmt = ws[cell_ref].number_format
    if fmt == "@":
        # text-formatted cell -- needs a literal fraction string, e.g. '1 1/8'
        if isinstance(bolt_od, str):
            ws[cell_ref] = bolt_od
        else:
            ws[cell_ref] = _decimal_to_fraction_string(bolt_od)
    else:
        if isinstance(bolt_od, str):
            ws[cell_ref] = _fraction_string_to_decimal(bolt_od)
        else:
            ws[cell_ref] = bolt_od


_FRACTIONS = {0.125: "1/8", 0.25: "1/4", 0.375: "3/8", 0.5: "1/2",
              0.625: "5/8", 0.75: "3/4", 0.875: "7/8"}


def _decimal_to_fraction_string(value):
    whole = int(value)
    frac = round(value - whole, 3)
    frac_str = _FRACTIONS.get(frac, "")
    if whole and frac_str:
        return f"{whole} {frac_str}"
    if whole:
        return str(whole)
    return frac_str or str(value)


def _fraction_string_to_decimal(s):
    s = s.strip().replace('"', "")
    parts = s.split()
    total = 0.0
    for p in parts:
        if "/" in p:
            n, d = p.split("/")
            total += float(n) / float(d)
        else:
            total += float(p)
    return total


def build_workbook(entries, template_path=TEMPLATE_PATH):
    """
    entries: list of dicts, each with:
      tab, pid, flange_no, tag, date, size, cls, bolt_od, boltcount,
      torqued_by, stages (4-tuple), tool_label, tool_serial, tool_cal,
      tool_row (default 39), system (default 'PH02'), comment (optional --
      used only for the corrections log, never written to the report)
    Returns an openpyxl Workbook with one sheet per entry, master/example
    sheets removed.
    """
    wb, imgs = _load_master(template_path)

    built_names = []
    for e in entries:
        boltcount = e["boltcount"]
        sheet_name = SHEET_BY_BOLTCOUNT.get(boltcount)
        if sheet_name is None or sheet_name not in wb.sheetnames:
            raise ValueError(f"No template sheet for bolt count {boltcount}")
        src = wb[sheet_name]
        ws = wb.copy_worksheet(src)
        # unique, Excel-safe tab name (<=31 chars, no duplicates)
        title = str(e.get("tab") or e.get("tag") or uuid.uuid4().hex[:8])[:31]
        base_title = title
        n = 1
        existing = {s.title for s in wb.worksheets}
        while title in existing:
            n += 1
            title = f"{base_title[:28]}-{n}"
        ws.title = title
        built_names.append(title)

        ws["B4"] = e.get("system", "PH02")
        ws["G4"] = e.get("date", "")
        ws["B6"] = e.get("pid", "")
        ws["B8"] = e.get("flange_no", "")
        ws["E8"] = e.get("tag", "")
        ws["B16"] = e.get("size", "")
        ws["D16"] = e.get("cls", "")
        _bolt_od_value(ws, "F16", e.get("bolt_od", ""))
        if e.get("gasket"):
            ws["H16"] = e["gasket"]
        ws["G18"] = "X"
        stages = e.get("stages") or (None, None, None, None)
        ws["B24"], ws["B25"], ws["B26"], ws["B27"] = stages
        ws["B28"] = None
        if e.get("lubricant_type"):
            ws["H23"] = e["lubricant_type"]
        if e.get("joint_material"):
            ws["H24"] = e["joint_material"]
        if e.get("washers"):
            ws["H25"] = e["washers"]
        if e.get("insul_kit"):
            ws["H26"] = e["insul_kit"]
        if e.get("bolt_material"):
            ws["H27"] = e["bolt_material"]
        torqued_by = e.get("torqued_by", "")
        if sheet_name == "4 Bolt":
            ws["A31"] = torqued_by
        else:
            ws["B31"] = torqued_by
        # Tools Required grid: clear the whole grid first so no leftover
        # template sample data bleeds through on a row this entry doesn't
        # use, then write the tool's serial/cal date to its capacity row
        # (see module docstring). Row 40's label is the only one ever
        # overwritten (an open/custom row); every other row keeps its
        # preprinted label.
        for r in range(32, 41):
            ws[f"H{r}"] = None
            ws[f"K{r}"] = None
        tool_row = e.get("tool_row", 39)
        if tool_row == 40 and e.get("tool_label"):
            ws[f"F{tool_row}"] = e["tool_label"]
        ws[f"H{tool_row}"] = e.get("tool_serial", "")
        ws[f"K{tool_row}"] = e.get("tool_cal", "")
        # Comments (A35) intentionally left untouched -- QC flags go to the
        # corrections log, never on the report itself.
        # Stranco Rep block: auto-filled per standing rule.
        ws["A57"] = STRANCO_REP_NAME
        ws["G57"] = None
        ws["J57"] = e.get("date", "")
        # QA/QC Rep block: left blank for a wet signature.
        for c in ["A59", "G59", "J59"]:
            ws[c] = None

        _add_images(ws, imgs, boltcount)
        _add_signature_image(ws)

    for name in MASTER_SHEETS:
        if name in wb.sheetnames:
            del wb[name]
    wb._sheets = [wb[n] for n in built_names]
    return wb


def workbook_to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def xlsx_bytes_to_pdf_bytes(xlsx_bytes, timeout=120):
    """Uses headless LibreOffice to convert the generated workbook to PDF,
    which is what reproduces the exact printable layout (fonts, borders,
    embedded images) rather than re-drawing it from scratch."""
    with tempfile.TemporaryDirectory() as tmp:
        xlsx_path = os.path.join(tmp, "report.xlsx")
        with open(xlsx_path, "wb") as f:
            f.write(xlsx_bytes)
        subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                "--convert-to", "pdf", "--outdir", tmp, xlsx_path,
            ],
            check=True, timeout=timeout,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        )
        pdf_path = os.path.join(tmp, "report.pdf")
        with open(pdf_path, "rb") as f:
            return f.read()
